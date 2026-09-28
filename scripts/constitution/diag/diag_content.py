# diag_content.py — 콘텐츠·데이터 파일 점검 (단계 S03·S06)
# 보는 것:
#   ① 수식($...$) 안의 제어 문자 — YAML·JSON 큰따옴표 이스케이프로 LaTeX가 손상된 흔적 (KR-06, R21-04안)
#      예: "\frac" → 폼피드+"rac", "\theta" → 탭+"heta", "\neq" → 줄바꿈+"eq", "\alpha" → 벨+"lpha"
#   ② 수식 밖의 이상한 제어 문자 (벨·백스페이스·폼피드 등)
#   ③ $ 개수가 홀수인 문자열 (수식 괄호 짝 불일치 의심)
#   ④ ID 형식 혼재 — 같은 키(예: concept_id)에 'H:12'·'H_12'·'[12...]'가 섞였는지 (KR-01)
# 지원 형식: .json .jsonl .yaml/.yml(PyYAML 있을 때) .csv .tsv .md .txt .xlsx(openpyxl 있을 때)
# 사용: python diag_content.py --path <폴더나 파일> [--path ...] --out <결과>/raw
from __future__ import annotations

import argparse
import csv
import io
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from _diagcommon import iter_files, md_table, setup_console, write_outputs

MATH = re.compile(r"\$\$(.+?)\$\$|\$([^$]+)\$", re.S)
CTRL_ANY = re.compile(r"[\x00-\x1f]")
CTRL_ODD = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")  # 탭·줄바꿈·CR 외의 제어 문자
LOST_CMD = re.compile(
    r"(\x0c(?=rac|orall)|\t(?=heta|imes|an\b|ext|au|riangle)|\n(?=eq|abla|u\b|ot)|"
    r"\x07(?=lpha|pprox|ngle)|\x08(?=eta|inom|ar|ig)|\x0b(?=ec|arepsilon|dots)|\r(?=ightarrow|ho|m))"
)
DEFAULT_ID_KEYS = (
    r"(?i)(^|_)(id|ids|code|codes|node|nodes|concept|standard|std)($|_)|코드|개념|노드|성취기준"
)


def shape(v: str) -> str:
    """ID의 모양: 한글→가, 숫자→9, 대문자→A, 소문자→a (구분자는 그대로) — 예: 'H:12' → 'A:99'"""
    out = []
    for ch in v:
        if "가" <= ch <= "힣":
            out.append("가")
        elif ch.isdigit():
            out.append("9")
        elif "A" <= ch <= "Z":
            out.append("A")
        elif "a" <= ch <= "z":
            out.append("a")
        else:
            out.append(ch)
    return "".join(out)


def separators(v: str) -> str:
    """ID 속 구분 기호만 남긴 모양: 'H:12'→':', 'H_13'→'_', '[12수학01-01]'→'[-]', 'J0101-K1'→'-'"""
    out = []
    for ch in v:
        if not (ch.isalnum() or "가" <= ch <= "힣") and (not out or out[-1] != ch):
            out.append(ch)
    return "".join(out) or "(없음)"


class Scanner:
    def __init__(self, id_keys: str):
        self.id_key = re.compile(id_keys)
        self.corrupt, self.ctrl, self.odd = [], [], []
        self.shapes = defaultdict(Counter)
        self.seps = defaultdict(Counter)  # 키별 구분자 모양 분포
        self.examples = {}
        self.strings = 0

    def check(self, s: str, where: str, key: str | None = None) -> None:
        self.strings += 1
        if key and self.id_key.search(key) and 0 < len(s) <= 40 and "\n" not in s:
            sh = shape(s)
            self.shapes[key][sh] += 1
            self.examples.setdefault((key, sh), s)
            sep = separators(s)
            self.seps[key][sep] += 1
            self.examples.setdefault((key, "sep:" + sep), s)
        for m in MATH.finditer(s):
            seg = m.group(1) or m.group(2) or ""
            if CTRL_ANY.search(seg):
                lost = LOST_CMD.search(seg)
                self.corrupt.append(
                    {
                        "where": where,
                        "text": repr(seg[:60]),
                        "guess": (
                            f"LaTeX 명령 손상 추정 ({repr(lost.group(0))} 뒤)"
                            if lost
                            else "제어 문자 포함"
                        ),
                    }
                )
        if CTRL_ODD.search(MATH.sub("", s)):
            self.ctrl.append({"where": where, "text": repr(s[:60])})
        if (s.count("$") - s.count("\\$")) % 2 == 1:
            self.odd.append({"where": where, "text": repr(s[:60])})

    def walk(self, v, where: str, key: str | None = None) -> None:
        if isinstance(v, str):
            self.check(v, where, key)
        elif isinstance(v, dict):
            for k, x in v.items():
                self.walk(x, f"{where}/{k}", str(k))
        elif isinstance(v, list):
            for i, x in enumerate(v):
                self.walk(x, f"{where}[{i}]", key)


def scan_file(sc: Scanner, p: Path, notes: list, max_mb: float) -> None:
    ext = p.suffix.lower()
    if p.stat().st_size > max_mb * 1024 * 1024:
        notes.append(f"{p}: {max_mb}MB 초과로 건너뜀")
        return
    try:
        if ext == ".xlsx":
            try:
                from openpyxl import load_workbook
            except ImportError:
                notes.append(f"{p}: openpyxl 없음 — xlsx 건너뜀")
                return
            wb = load_workbook(p, read_only=True, data_only=True)
            for ws in wb.worksheets:
                header = None
                for i, row in enumerate(ws.iter_rows(values_only=True), 1):
                    if header is None:
                        header = [str(c) if c is not None else "" for c in row]
                        continue
                    for j, c in enumerate(row):
                        if isinstance(c, str):
                            sc.check(
                                c,
                                f"{p.name}!{ws.title}!R{i}C{j + 1}",
                                header[j] if j < len(header) else None,
                            )
            wb.close()
            return
        text = p.read_text(encoding="utf-8-sig")
        if ext == ".json":
            sc.walk(json.loads(text), p.name)
        elif ext == ".jsonl":
            for i, line in enumerate(text.splitlines(), 1):
                if line.strip():
                    sc.walk(json.loads(line), f"{p.name}:{i}")
        elif ext in (".yaml", ".yml"):
            try:
                import yaml
            except ImportError:
                notes.append(f"{p}: PyYAML 없음 — YAML 건너뜀")
                return
            for k, doc in enumerate(yaml.safe_load_all(text)):
                sc.walk(doc, f"{p.name}#{k}")
        elif ext in (".csv", ".tsv"):
            rows = list(csv.reader(io.StringIO(text), delimiter="\t" if ext == ".tsv" else ","))
            header = rows[0] if rows else []
            for i, row in enumerate(rows[1:], 2):
                for j, c in enumerate(row):
                    sc.check(c, f"{p.name}:R{i}C{j + 1}", header[j] if j < len(header) else None)
        else:  # .md .txt — 파일 전체를 한 문자열로
            sc.check(text, p.name)
    except (UnicodeDecodeError, json.JSONDecodeError, ValueError, OSError) as e:
        notes.append(f"{p}: 읽기 실패 — {type(e).__name__}: {str(e)[:120]}")


def main() -> int:
    setup_console()
    ap = argparse.ArgumentParser(description="콘텐츠·데이터 파일 점검")
    ap.add_argument(
        "--path",
        action="append",
        required=True,
        type=Path,
        help="점검할 폴더나 파일 (여러 번 지정 가능)",
    )
    ap.add_argument("--out", required=True, type=Path)
    ap.add_argument("--max-mb", type=float, default=200)
    ap.add_argument("--id-keys", default=DEFAULT_ID_KEYS, help="ID로 보고 형식을 셀 키 이름 정규식")
    args = ap.parse_args()
    sc, notes, files = Scanner(args.id_keys), [], []
    exts = {".json", ".jsonl", ".yaml", ".yml", ".csv", ".tsv", ".md", ".txt", ".xlsx"}
    for base in args.path:
        targets = (
            [base] if base.is_file() else list(iter_files(base, exts)) if base.is_dir() else []
        )
        if not targets:
            notes.append(f"{base}: 경로가 없거나 대상 파일 없음")
        for p in targets:
            files.append(str(p))
            scan_file(sc, p, notes, args.max_mb)

    mixed = {}
    for (
        key,
        cnt,
    ) in sc.seps.items():  # 같은 키에 구분자 모양이 둘 이상 = 형식 혼재 (조인이 조용히 깨짐)
        if len(cnt) > 1:
            mixed[key] = [
                {"shape": s, "count": n, "example": sc.examples[(key, "sep:" + s)]}
                for s, n in cnt.most_common(8)
            ]
    data = {
        "paths": [str(p) for p in args.path],
        "files": len(files),
        "strings": sc.strings,
        "summary": {
            "corrupted_math": len(sc.corrupt),
            "odd_control_chars": len(sc.ctrl),
            "odd_dollar": len(sc.odd),
            "id_keys_mixed": len(mixed),
        },
        "corrupted_math": sc.corrupt,
        "odd_control_chars": sc.ctrl,
        "odd_dollar": sc.odd,
        "id_shapes_mixed": mixed,
        "id_separators_all": {k: dict(v.most_common(10)) for k, v in sc.seps.items()},
        "id_shapes_all": {k: dict(v.most_common(10)) for k, v in sc.shapes.items()},
        "notes": notes,
    }
    s = data["summary"]
    md = [
        "# 콘텐츠·데이터 파일 점검\n",
        f"대상: {', '.join(f'`{p}`' for p in args.path)} · 파일 {len(files)}개 · 문자열 {sc.strings:,}개\n",
        md_table(
            ["항목", "값"],
            [
                ["⛔ 수식 안 제어 문자 (LaTeX 이스케이프 손상, KR-06)", s["corrupted_math"]],
                ["수식 밖 이상 제어 문자", s["odd_control_chars"]],
                ["$ 개수가 홀수인 문자열", s["odd_dollar"]],
                ["ID 형식이 섞인 키 (KR-01)", s["id_keys_mixed"]],
            ],
        ),
        "\n## 손상된 수식\n",
        md_table(
            ["위치", "수식(보이지 않는 문자 표시)", "추정"],
            [[c["where"], c["text"], c["guess"]] for c in sc.corrupt],
            50,
        ),
        "\n## ID 형식 혼재\n",
        md_table(
            ["키", "구분자 모양", "개수", "예시"],
            [[k, x["shape"], x["count"], x["example"]] for k, xs in mixed.items() for x in xs],
            60,
        ),
        "\n## 수식 밖 제어 문자\n",
        md_table(["위치", "문자열"], [[c["where"], c["text"]] for c in sc.ctrl], 30),
        "\n## $ 짝이 안 맞는 문자열\n",
        md_table(["위치", "문자열"], [[c["where"], c["text"]] for c in sc.odd], 30),
        "\n## 참고\n" + ("\n".join(f"- {n}" for n in notes) or "- 없음"),
    ]
    write_outputs(args.out, "diag_content", data, "\n".join(md) + "\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
